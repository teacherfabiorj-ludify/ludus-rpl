"""
LUDUS — helpers compartilhados pelos build_*.py das planilhas.

Lê core/core.json — o dump de core/*.js gerado por `node core/export.js`.
NENHUM script de planilha pode redigitar uma regra, um termo ou um número:
a verdade vem do JSON, igual aos livros.

Fórmulas são escritas com VÍRGULA (padrão Excel). O Google traduz para ponto e
vírgula sozinho na conversão do .xlsx — testado. Nunca escreva ponto e vírgula
aqui: isso quebra no Excel e no LibreOffice.
"""

import json
import os

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, "..", ".."))
OUT = os.path.abspath(os.path.join(HERE, ".."))


def core():
    p = os.path.join(REPO, "core", "core.json")
    if not os.path.exists(p):
        raise SystemExit(
            "core/core.json não existe.\n"
            "Rode primeiro:  node core/export.js"
        )
    with open(p, encoding="utf-8") as f:
        return json.load(f)


# --- paleta: a mesma dos livros -------------------------------------------
INK = "14130F"
ACCENT = "2A78D6"
BRAND = "D2691E"
GOOD = "1F7A4D"
WARN = "C98510"
CRIT = "C4372F"

FILL_HEAD = PatternFill("solid", fgColor=INK)
FILL_TYPE = PatternFill("solid", fgColor="FFF4D6")   # amarelo: você digita
FILL_CALC = PatternFill("solid", fgColor="EAF2FD")   # azul: calculado
FILL_PRIV = PatternFill("solid", fgColor="FBEDEC")   # vermelho claro: privado
FILL_NOTE = PatternFill("solid", fgColor="F4F3F0")

F_HEAD = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
F_TITLE = Font(name="Calibri", size=16, bold=True, color=INK)
F_SUB = Font(name="Calibri", size=10, color="6E6A61")
F_BODY = Font(name="Calibri", size=10, color=INK)
F_LABEL = Font(name="Calibri", size=10, bold=True, color=INK)
F_SMALL = Font(name="Calibri", size=9, color="6E6A61")

_thin = Side(style="thin", color="D8D6D0")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

WRAP = Alignment(vertical="top", wrap_text=True)
WRAP_C = Alignment(vertical="center", horizontal="center", wrap_text=True)


def widths(ws, spec):
    """spec: {'A': 22, 'B': 10, ...}"""
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def title(ws, row, text, sub=None):
    ws.cell(row, 1, text).font = F_TITLE
    if sub:
        ws.cell(row + 1, 1, sub).font = F_SUB
        return row + 3
    return row + 2


def header_row(ws, row, labels, start_col=1):
    for i, lab in enumerate(labels):
        c = ws.cell(row, start_col + i, lab)
        c.fill, c.font, c.alignment, c.border = FILL_HEAD, F_HEAD, WRAP_C, BORDER
    ws.row_dimensions[row].height = 28
    return row + 1


def note(ws, row, text, font=None):
    c = ws.cell(row, 1, text)
    c.font = font or F_SMALL
    c.alignment = WRAP
    return row + 1


def paint(ws, row, cols, fill):
    for col in cols:
        c = ws[f"{col}{row}"]
        c.fill = fill
        c.border = BORDER
        c.alignment = WRAP
        c.font = F_BODY


def freeze(ws, cell):
    ws.freeze_panes = cell


def leia_me(wb, lines, heading, sub):
    ws = wb.create_sheet("LEIA-ME", 0)
    widths(ws, {"A": 120})
    r = title(ws, 2, heading, sub)
    for line in lines:
        if line == "":
            r += 1
            continue
        if line.startswith("## "):
            c = ws.cell(r, 1, line[3:])
            c.font = F_LABEL
            r += 1
            continue
        c = ws.cell(r, 1, line)
        c.font = F_BODY
        c.alignment = WRAP
        r += 1
    return ws
