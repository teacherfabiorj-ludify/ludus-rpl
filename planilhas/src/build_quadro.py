"""
LUDUS — Class Board  (the whole class sees it · READ ONLY)

    node core/export.js
    python3 planilhas/src/build_quadro.py

One file, two tabs. It is the middle link of the chain:

    Painel (teacher's)  →  THIS FILE  →  each student's character sheet

It exists for one reason: IMPORTRANGE is authorised per WHOLE FILE, not per
range. If a student's sheet pointed straight at the Painel, they would only
have to change the range inside the formula to read the whole class's marks,
attempts and alerts. Pointing here, the most they can reach is what they could
already open.

⚠ 19/09/2026 — this file is now in ENGLISH, and its tab is called BOARD. The
students read it every session; it was in Portuguese, in an English course.
The teacher's Painel stays in Portuguese — it never leaves the teacher's Drive.
"""

import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from _common import (ACCENT, BORDER, BOARD_COL, BOARD_ROW, FILL_TYPE, F_BODY,
                     F_LABEL, F_SMALL, INK, OUT, PAINEL_PUBLIC, core, leia_me,
                     widths)

C = core()
GAME = C["brand"]["GAME_NAME"]
HOUSE = C["brand"]["HOUSE"]
VERSION = C["brand"]["VERSION"]

IMPORT_ROW = BOARD_ROW  # where the IMPORTRANGE lives — see _common.board_public()
ROWS = 4

FILL_TITLE = PatternFill("solid", fgColor=INK)
F_MAST = Font(name="Calibri", size=20, bold=True, color="FFFFFF")
F_MASTSUB = Font(name="Calibri", size=10, color="E8E6E0")
LEFT = Alignment(vertical="center", wrap_text=True)


def sheet_board(wb):
    ws = wb.create_sheet("BOARD")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = ACCENT
    widths(ws, {"A": 2, "B": 18, "C": 10, "D": 34, "E": 9, "F": 32, "G": 14,
                "H": 26, "I": 10, "J": 48, "K": 2})

    for col in range(2, 11):
        ws.cell(2, col).fill = FILL_TITLE
        ws.cell(3, col).fill = FILL_TITLE
    ws.cell(2, 2, f"{GAME} · CLASS BOARD").font = F_MAST
    ws.cell(3, 2, f"{HOUSE}  ·  {VERSION}  ·  everything here comes from the teacher's panel. "
                  "Nobody types in this file.").font = F_MASTSUB
    ws.row_dimensions[2].height = 30
    ws.row_dimensions[3].height = 17

    c = ws.cell(5, 2, "PASTE THE PANEL LINK HERE →")
    c.font, c.alignment = F_LABEL, LEFT
    v = ws.cell(5, 3, "paste the Painel-Turma link here")
    v.fill, v.font, v.border, v.alignment = FILL_TYPE, F_BODY, BORDER, LEFT
    ws.merge_cells(start_row=5, start_column=3, end_row=5, end_column=7)

    f = ws.cell(IMPORT_ROW, BOARD_COL, f'=IMPORTRANGE($C$5,"{PAINEL_PUBLIC}")')
    f.font, f.border = F_BODY, BORDER

    n = IMPORT_ROW + ROWS + 4
    def note_(row, text, font=None):
        cell = ws.cell(row, 2, text)
        cell.font, cell.alignment = font or F_SMALL, LEFT
        return row + 1

    n = note_(n, f"FIRST TIME: CELL B{IMPORT_ROW} WILL SHOW #REF!", F_LABEL)
    for line in [
        f"Click cell B{IMPORT_ROW} and press the Allow access button. That is what authorises "
        "the IMPORTRANGE. It happens once, and never asks again.",
        "Before that, paste the panel link into the yellow cell C5.",
        "",
        "## What this file does NOT have",
        "No marks. No attempts. No alerts. Nothing a student may not see.",
        "That is why it can be shared with the whole class — and why the character sheets read "
        "this file, not the teacher's panel.",
        "",
        "## How to share it",
        "Share → General access → Anyone with the link → VIEWER. Never Editor: nobody needs to "
        "type in here.",
        "In Classroom, post it as MATERIAL under the Quick Reference topic.",
    ]:
        if line.startswith("## "):
            n = note_(n, line[3:], F_LABEL)
        else:
            n = note_(n, line, F_BODY if line else F_SMALL)
    return ws


def build():
    wb = Workbook()
    wb.remove(wb.active)
    sheet_board(wb)

    ws = leia_me(wb, [
        f"{VERSION}  ·  generated from core/core.json by planilhas/src/build_quadro.py",
        "",
        "## What it is for",
        "This is the middle file. The teacher's panel has everything — marks, attempts, alerts. The student's sheet has the character. Between the two sits this one, read-only, carrying only what can be public.",
        "",
        "## Setting it up, once",
        "1. Upload it to Drive, in the class folder (not the GM folder), and open it.",
        "2. File → Save as Google Sheets.",
        "3. Paste the panel link into the yellow cell C5 of the BOARD tab.",
        "4. Click cell B6 (it will show #REF!) and press Allow access.",
        "5. Share → Anyone with the link → Viewer. Copy the link: that is what goes into the character sheets and into Classroom.",
        "",
        "## Why not link the character sheets straight to the panel",
        "IMPORTRANGE is authorised per whole file, not per range. Once authorised inside a student's sheet, they can change A1:I5 to A1:Z200 — it is a text field, in their own spreadsheet — and read the entire panel. Pointing at this file, the most they can reach is the board they could already open.",
        "",
        "## Language",
        "This file is in English because the class reads it. The teacher's panel stays in Portuguese: it is the teacher's private tool.",
        "",
        "## Refresh",
        "IMPORTRANGE is not instant: Sheets recalculates every few minutes, or when the file is opened. For the real use — you prepare beforehand, they open it during the session — that is invisible.",
    ], f"Class Board · {GAME}", "The whole class sees it. Nobody edits it.",
        name="READ ME", pos=None)
    ws.sheet_view.showGridLines = False

    out = os.path.join(OUT, "Class-Board-LUDUS.xlsx")
    wb.save(out)
    print("written:", out, "· tabs:", wb.sheetnames)


if __name__ == "__main__":
    build()
