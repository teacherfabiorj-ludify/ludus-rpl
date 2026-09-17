"""
LUDUS — ★ as 72 unidades, do catálogo para o core

    python3 planilhas/src/export_units.py     → core/units.json

Roda ANTES de `node core/export.js`. Até 16/09 as 72 unidades só existiam
dentro do Catalogo-Etiquetas-Evolve.xlsx, e o Painel as lia direto de lá —
o que significava que o Core Book não tinha como imprimir a mesma tabela sem
redigitá-la. Agora o catálogo despeja aqui, e livro e planilha leem o mesmo
arquivo.

⚠ ESTA É A CAMADA DA TRILHA. É o único lugar do repositório onde existe a
palavra "Evolve" seguida de um número de unidade. Trocar de material didático
é regerar este arquivo de outro catálogo.
"""

import json
import os

import openpyxl

from _common import OUT, REPO

CATALOGUE = os.path.join(OUT, "Catalogo-Etiquetas-Evolve.xlsx")
TARGET = os.path.join(REPO, "core", "units.json")



# ---------------------------------------------------------------------------
# ⚠ O catálogo foi escrito em português de planejamento, e os livros são em
# inglês. Sem esta normalização a tabela oficial sai com "Simple present
# (hábitos)" dentro de um livro inglês — que é exatamente o tipo de mistura que
# a regra de idioma proíbe. Ordem importa: específicos primeiro, genéricos no
# fim.
# ---------------------------------------------------------------------------
NORMALISE = [
    ("(afirm., neg. e pergunta)", "(affirmative, negative and question forms)"),
    ("(afirm. e pergunta)", "(affirmative and question forms)"),
    ("(habilidade e possibilidade)", "(ability and possibility)"),
    ("pergunta sim/não", "yes/no questions"),
    ("(hábitos)", "(habits)"),
    ("(planos futuros)", "(future plans)"),
    ("(experiências)", "(experiences)"),
    ("(decisão súbita)", "(sudden decisions)"),
    ("(presente e futuro)", "(present and future)"),
    ("(presente e passado)", "(present and past)"),
    ("(presente, futuro e passado)", "(present, future and past)"),
    ("(futuro)", "(future)"),
    ("contraste simple present x continuous", "simple present vs continuous"),
    ("contraste com", "contrast with"),
    ("Modais de necessidade", "Modals of necessity"),
    ("modais de proibição e permissão", "modals of prohibition and permission"),
    ("Variações de", "Variations on"),
    ("modais", "modals"),
    ("após", "after"),
    (" em relative clauses", " in relative clauses"),
    (" com ", " with "),
    (" e ", " and "),
]


def to_english(text):
    t = text or ""
    for a, b in NORMALISE:
        t = t.replace(a, b)
    return t

def split_lessons(grammar):
    """O catálogo traz a gramática por UNIDADE. A divisão em lição 1 / lição 2
    é inferida cortando no primeiro ponto e vírgula — e é marcada como tal."""
    g = to_english((grammar or "").strip())
    if ";" in g:
        i = g.index(";")
        return g[:i].strip(), g[i + 1:].strip(), "inferred"
    return g, "⚠ check the book — the catalogue gives only one topic", "⚠ CHECK"


def build():
    wb = openpyxl.load_workbook(CATALOGUE)
    ws = wb["MAPA"]
    units = []
    for r in range(5, ws.max_row + 1):
        ev = ws.cell(r, 2).value
        if ev is None:
            continue
        l1, l2, conf = split_lessons(ws.cell(r, 5).value)
        units.append({
            "key": f"E{int(ev)}-U{int(ws.cell(r, 3).value)}",
            "cefr": ws.cell(r, 1).value,
            "level": int(ev),
            "unit": int(ws.cell(r, 3).value),
            "title": ws.cell(r, 4).value or "",
            "lesson1": l1,
            "lesson2": l2,
            "action1": ws.cell(r, 6).value or "—",
            "action2": ws.cell(r, 7).value or "—",
            "domain": ws.cell(r, 8).value or "",
            "checked": conf,
        })

    payload = {
        "_generated_from": "Catalogo-Etiquetas-Evolve.xlsx · MAPA",
        "_warning": "GENERATED FILE — do not edit. Run planilhas/src/export_units.py",
        "_trailLayer": True,
        "coursebook": "Evolve",
        "count": len(units),
        "units": units,
    }
    with open(TARGET, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=1)
    inferred = sum(1 for u in units if u["checked"] != "inferred")
    print(f"escrito: {TARGET} · {len(units)} unidades · {inferred} a conferir no livro")


if __name__ == "__main__":
    build()
